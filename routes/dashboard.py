"""
Comprehensive Dashboard Routes for viewing registered teams
in the Inter-Dept Sports Meet (IIT Roorkee).
"""

import io
import csv
from datetime import datetime
from typing import Any, Optional
from fastapi import APIRouter, Request, Query, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
import os

from database.config import get_client, ping_db, DATABASE_NAME

router = APIRouter(tags=["Teams Dashboard"])

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
template_candidates = [
    os.path.join(BASE_DIR, "templates"),
    os.path.join(os.getcwd(), "templates"),
    os.path.abspath("templates"),
]
templates_dir = next((p for p in template_candidates if os.path.isdir(p)), os.path.join(BASE_DIR, "templates"))
templates = Jinja2Templates(directory=templates_dir)

# Competing Departments Definition
DEPARTMENT_CONFIG: dict[str, dict[str, str]] = {
    "amsc": {
        "code": "amsc",
        "abbr": "AMSC",
        "name": "Applied Mathematics & Scientific Computing",
        "badge_color": "#38bdf8",
        "bg_color": "rgba(56, 189, 248, 0.12)",
        "border_color": "rgba(56, 189, 248, 0.35)",
    },
    "ppe": {
        "code": "ppe",
        "abbr": "PPE",
        "name": "Polymer and Process Engineering",
        "badge_color": "#a855f7",
        "bg_color": "rgba(168, 85, 247, 0.12)",
        "border_color": "rgba(168, 85, 247, 0.35)",
    },
    "ppt": {
        "code": "ppt",
        "abbr": "PPT",
        "name": "Paper and Pulp Technology",
        "badge_color": "#10b981",
        "bg_color": "rgba(16, 185, 129, 0.12)",
        "border_color": "rgba(16, 185, 129, 0.35)",
    },
}

# Subdomain / alias mapping
DEPT_ALIASES = {
    "amsc": "amsc",
    "pe": "ppe",
    "ppe": "ppe",
    "pt": "ppt",
    "ppt": "ppt",
}

# 16 Canonical Tournament Sports
SPORTS_CONFIG: dict[str, dict[str, Any]] = {
    "cricket_boys": {
        "id": "cricket_boys",
        "name": "Cricket",
        "display_name": "Cricket (boys)",
        "category": "Boys",
        "icon": "🏏",
        "main_count": 11,
        "sub_max": 4,
        "formula": "11 + 4",
        "squad_notation": "11 Playing + up to 4 Subs",
        "description": "11 Playing (1 Captain + 10 Main) + up to 4 Substitutes. Total squad: 11 to 15 Boys.",
    },
    "box_cricket_girls": {
        "id": "box_cricket_girls",
        "name": "Box Cricket",
        "display_name": "Box cricket (girls)",
        "category": "Girls",
        "icon": "🏏",
        "main_count": 5,
        "sub_max": 3,
        "formula": "5 + 3",
        "squad_notation": "5 Playing + up to 3 Subs",
        "description": "5 Playing (1 Captain + 4 Main) + up to 3 Substitutes. Total squad: 5 to 8 Girls.",
    },
    "football_boys": {
        "id": "football_boys",
        "name": "Football",
        "display_name": "Football (boys)",
        "category": "Boys",
        "icon": "⚽",
        "main_count": 7,
        "sub_max": 3,
        "formula": "7 + 3",
        "squad_notation": "7 Playing + up to 3 Subs",
        "description": "7 Playing (1 Captain + 6 Main) + up to 3 Substitutes. Total squad: 7 to 10 Boys.",
    },
    "football_girls": {
        "id": "football_girls",
        "name": "Football",
        "display_name": "Football (girls)",
        "category": "Girls",
        "icon": "⚽",
        "main_count": 3,
        "sub_max": 2,
        "formula": "3 + 2",
        "squad_notation": "3 Playing + up to 2 Subs",
        "description": "3 Playing (1 Captain + 2 Main) + up to 2 Substitutes. Total squad: 3 to 5 Girls.",
    },
    "volleyball_boys": {
        "id": "volleyball_boys",
        "name": "Volleyball",
        "display_name": "Volleyball (boys)",
        "category": "Boys",
        "icon": "🏐",
        "main_count": 6,
        "sub_max": 2,
        "formula": "6 + 2",
        "squad_notation": "6 Playing + up to 2 Subs",
        "description": "6 Playing (1 Captain + 5 Main) + up to 2 Substitutes. Total squad: 6 to 8 Boys.",
    },
    "badminton_mix": {
        "id": "badminton_mix",
        "name": "Badminton",
        "display_name": "Badminton (mix)",
        "category": "Mixed",
        "icon": "🏸",
        "main_count": 4,
        "sub_max": 2,
        "formula": "2B + 2G + up to 2",
        "squad_notation": "4 Playing (2B + 2G) + up to 2 Subs",
        "description": "4 Playing (exactly 2 Boys & 2 Girls) + up to 2 Substitutes (max 1B, 1G).",
    },
    "tt_boys": {
        "id": "tt_boys",
        "name": "Table Tennis",
        "display_name": "TT (boys)",
        "category": "Boys",
        "icon": "🏓",
        "main_count": 3,
        "sub_max": 2,
        "formula": "3 + 2",
        "squad_notation": "3 Playing + up to 2 Subs",
        "description": "3 Playing (1 Captain + 2 Main) + up to 2 Substitutes. Total squad: 3 to 5 Boys.",
    },
    "carrom_mix": {
        "id": "carrom_mix",
        "name": "Carrom",
        "display_name": "Carrom (mix)",
        "category": "Mixed",
        "icon": "🎯",
        "main_count": 2,
        "sub_max": 2,
        "formula": "1B + 1G + up to 2",
        "squad_notation": "2 Playing (1B + 1G) + up to 2 Subs",
        "description": "2 Playing (1 Boy & 1 Girl) + up to 2 Substitutes (max 1B, 1G). Total: 2 to 4.",
    },
    "tow_boys": {
        "id": "tow_boys",
        "name": "Tug of War",
        "display_name": "TOW (boys)",
        "category": "Boys",
        "icon": "🪢",
        "main_count": 7,
        "sub_max": 2,
        "formula": "7 + 2",
        "squad_notation": "7 Playing + up to 2 Subs",
        "description": "7 Playing (1 Captain + 6 Main) + up to 2 Substitutes. Total squad: 7 to 9 Boys.",
    },
    "tow_girls": {
        "id": "tow_girls",
        "name": "Tug of War",
        "display_name": "TOW (girls)",
        "category": "Girls",
        "icon": "🪢",
        "main_count": 5,
        "sub_max": 2,
        "formula": "5 + 2",
        "squad_notation": "5 Playing + up to 2 Subs",
        "description": "5 Playing (1 Captain + 4 Main) + up to 2 Substitutes. Total squad: 5 to 7 Girls.",
    },
    "shotput_boys": {
        "id": "shotput_boys",
        "name": "Shot Put",
        "display_name": "Shotput (boys)",
        "category": "Boys",
        "icon": "☄️",
        "main_count": 2,
        "sub_max": 1,
        "formula": "2 + 1",
        "squad_notation": "2 Playing + up to 1 Sub",
        "description": "2 Playing (1 Captain + 1 Main) + up to 1 Substitute. Total: 2 to 3 Boys.",
    },
    "shotput_girls": {
        "id": "shotput_girls",
        "name": "Shot Put",
        "display_name": "Shotput (girls)",
        "category": "Girls",
        "icon": "☄️",
        "main_count": 1,
        "sub_max": 1,
        "formula": "1 + 1",
        "squad_notation": "1 Playing + up to 1 Sub",
        "description": "1 Playing (Captain) + up to 1 Substitute. Total: 1 to 2 Girls.",
    },
    "discus_boys": {
        "id": "discus_boys",
        "name": "Discus Throw",
        "display_name": "Discus throw (boys)",
        "category": "Boys",
        "icon": "💿",
        "main_count": 2,
        "sub_max": 1,
        "formula": "2 + 1",
        "squad_notation": "2 Playing + up to 1 Sub",
        "description": "2 Playing (1 Captain + 1 Main) + up to 1 Substitute. Total: 2 to 3 Boys.",
    },
    "discus_girls": {
        "id": "discus_girls",
        "name": "Discus Throw",
        "display_name": "Discus throw (girls)",
        "category": "Girls",
        "icon": "💿",
        "main_count": 1,
        "sub_max": 1,
        "formula": "1 + 1",
        "squad_notation": "1 Playing + up to 1 Sub",
        "description": "1 Playing (Captain) + up to 1 Substitute. Total: 1 to 2 Girls.",
    },
    "long_jump_boys": {
        "id": "long_jump_boys",
        "name": "Long Jump",
        "display_name": "Long jump (boys)",
        "category": "Boys",
        "icon": "🏃",
        "main_count": 2,
        "sub_max": 1,
        "formula": "2 + 1",
        "squad_notation": "2 Playing + up to 1 Sub",
        "description": "2 Playing (1 Captain + 1 Main) + up to 1 Substitute. Total: 2 to 3 Boys.",
    },
    "long_jump_girls": {
        "id": "long_jump_girls",
        "name": "Long Jump",
        "display_name": "Long jump (girls)",
        "category": "Girls",
        "icon": "🏃‍♀️",
        "main_count": 1,
        "sub_max": 1,
        "formula": "1 + 1",
        "squad_notation": "1 Playing + up to 1 Sub",
        "description": "1 Playing (Captain) + up to 1 Substitute. Total: 1 to 2 Girls.",
    },
}


def _format_date(dt_val: Any) -> str:
    """Format datetime value into readable string."""
    if not dt_val:
        return ""
    if isinstance(dt_val, str):
        try:
            # Parse ISO formatted string
            clean_str = dt_val.replace("Z", "+00:00")
            dt = datetime.fromisoformat(clean_str)
            return dt.strftime("%b %d, %Y · %I:%M %p")
        except Exception:
            return dt_val[:19].replace("T", " ")
    if isinstance(dt_val, datetime):
        return dt_val.strftime("%b %d, %Y · %I:%M %p")
    return str(dt_val)


async def fetch_all_teams() -> tuple[list[dict[str, Any]], Optional[str]]:
    """
    Fetch all registered teams across all department collections from MongoDB Atlas.
    Returns (teams_list, error_message).
    """
    all_teams: list[dict[str, Any]] = []
    error_msg: Optional[str] = None
    seen_team_ids: set[str] = set()

    try:
        client = get_client()
        db = client[DATABASE_NAME]

        # Discover department collections
        cols = await db.list_collection_names()
        target_depts = [c for c in cols if c not in ["users", "system.views", "system.profile"]]
        if not target_depts:
            target_depts = ["amsc", "ppe", "ppt"]

        for dept_raw in target_depts:
            canonical_dept = DEPT_ALIASES.get(dept_raw.lower(), dept_raw.lower())
            dept_info = DEPARTMENT_CONFIG.get(canonical_dept, {
                "code": canonical_dept,
                "abbr": canonical_dept.upper(),
                "name": canonical_dept.upper(),
                "badge_color": "#06b6d4",
                "bg_color": "rgba(6, 182, 212, 0.12)",
                "border_color": "rgba(6, 182, 212, 0.35)",
            })

            col = db[dept_raw]
            docs = await col.find({}).to_list(length=None)

            for doc in docs:
                doc_id = str(doc.get("_id", ""))

                # Case 1: Primary events document containing all sports
                if doc_id == "events":
                    for sport_id, sport_conf in SPORTS_CONFIG.items():
                        val = doc.get(sport_id)
                        items = val if isinstance(val, list) else ([val] if isinstance(val, dict) else [])
                        for t in items:
                            if isinstance(t, dict) and (t.get("captain") or t.get("created_by") or t.get("status")):
                                team = _normalize_team(t, sport_id, sport_conf, dept_info)
                                if team["team_id"] not in seen_team_ids:
                                    seen_team_ids.add(team["team_id"])
                                    all_teams.append(team)

                # Case 2: Individual sport document
                elif doc_id in SPORTS_CONFIG:
                    sport_id = doc_id
                    sport_conf = SPORTS_CONFIG[sport_id]
                    teams_list = doc.get("teams") or doc.get(sport_id) or []
                    if isinstance(teams_list, list):
                        for t in teams_list:
                            if isinstance(t, dict) and (t.get("captain") or t.get("created_by") or t.get("status")):
                                team = _normalize_team(t, sport_id, sport_conf, dept_info)
                                if team["team_id"] not in seen_team_ids:
                                    seen_team_ids.add(team["team_id"])
                                    all_teams.append(team)

    except Exception as exc:
        error_msg = f"MongoDB Atlas Connection: {str(exc)}"

    # Sort teams by sport name and department abbreviation
    all_teams.sort(key=lambda t: (t.get("sport_name", ""), t.get("department", "")))
    return all_teams, error_msg


def _normalize_team(
    raw: dict[str, Any],
    sport_id: str,
    sport_conf: dict[str, Any],
    dept_info: dict[str, Any],
) -> dict[str, Any]:
    """Ensure consistent structure, computed counts, and clean fields for every team."""
    team = dict(raw)

    captain = team.get("captain") or {}
    captain_email = (captain.get("email") or "").strip().lower()

    # Generate stable team ID if missing
    if not team.get("team_id"):
        team["team_id"] = f"{dept_info['code']}_{sport_id}_{captain_email or 'team'}"

    # Standardize department details
    team["department"] = dept_info["abbr"]
    team["department_name"] = dept_info["name"]
    team["department_code"] = dept_info["code"]
    team["dept_badge_color"] = dept_info.get("badge_color", "#06b6d4")
    team["dept_bg_color"] = dept_info.get("bg_color", "rgba(6, 182, 212, 0.12)")
    team["dept_border_color"] = dept_info.get("border_color", "rgba(6, 182, 212, 0.35)")

    # Standardize sport details
    team["sport_id"] = sport_id
    team["sport_name"] = sport_conf.get("display_name", sport_conf["name"])
    team["sport_clean_name"] = sport_conf["name"]
    team["sport_icon"] = sport_conf.get("icon", "🏅")
    team["category"] = sport_conf.get("category", "Open")
    team["formula"] = sport_conf.get("formula", "")
    team["squad_notation"] = sport_conf.get("squad_notation", "")

    # Captain normalization
    captain_record = {
        "name": (captain.get("name") or "Unnamed Captain").strip(),
        "email": (captain.get("email") or "").strip().lower(),
        "student_id": (captain.get("student_id") or "").strip(),
        "gender": captain.get("gender") or "Not Specified",
        "role": "Captain",
        "is_captain": True,
    }
    team["captain"] = captain_record

    # Main players normalization
    main_players = []
    for idx, p in enumerate(team.get("main_players") or [], start=1):
        if isinstance(p, dict):
            main_players.append({
                "name": (p.get("name") or f"Player {idx}").strip(),
                "email": (p.get("email") or "").strip().lower(),
                "student_id": (p.get("student_id") or "").strip(),
                "gender": p.get("gender") or "Not Specified",
                "role": p.get("role") or f"Main Player {idx}",
                "is_captain": False,
            })
    team["main_players"] = main_players

    # Substitutes normalization
    substitutes = []
    for idx, s in enumerate(team.get("substitutes") or [], start=1):
        if isinstance(s, dict):
            substitutes.append({
                "name": (s.get("name") or f"Sub {idx}").strip(),
                "email": (s.get("email") or "").strip().lower(),
                "student_id": (s.get("student_id") or "").strip(),
                "gender": s.get("gender") or "Not Specified",
                "role": s.get("role") or f"Substitute {idx}",
                "is_captain": False,
            })
    team["substitutes"] = substitutes

    # Total strength calculation
    playing_count = 1 + len(main_players)  # Captain + Main Players
    sub_count = len(substitutes)
    total_strength = playing_count + sub_count

    team["total_playing"] = team.get("total_playing") or playing_count
    team["total_substitutes"] = team.get("total_substitutes") or sub_count
    team["total_strength"] = team.get("total_strength") or total_strength

    # All members array for fast search / iteration
    team["all_members"] = [captain_record] + main_players + substitutes

    # Status & Timestamps
    team["status"] = team.get("status") or "Registered"
    team["created_at_formatted"] = _format_date(team.get("created_at"))
    team["updated_at_formatted"] = _format_date(team.get("updated_at") or team.get("created_at"))

    # Friendly team display name
    team["display_title"] = f"{team['department']} - {team['sport_name']}"

    return team


def compute_dashboard_stats(teams: list[dict[str, Any]]) -> dict[str, Any]:
    """Calculate aggregate statistics across all registered teams."""
    unique_athletes = set()
    dept_counts: dict[str, int] = {k: 0 for k in DEPARTMENT_CONFIG}
    sport_counts: dict[str, int] = {k: 0 for k in SPORTS_CONFIG}
    category_counts: dict[str, int] = {"Boys": 0, "Girls": 0, "Mixed": 0}
    gender_counts: dict[str, int] = {"Boy": 0, "Girl": 0, "Other": 0}

    for t in teams:
        dept_code = t.get("department_code")
        if dept_code in dept_counts:
            dept_counts[dept_code] += 1

        sport_id = t.get("sport_id")
        if sport_id in sport_counts:
            sport_counts[sport_id] += 1

        cat = t.get("category", "Boys")
        category_counts[cat] = category_counts.get(cat, 0) + 1

        for member in t.get("all_members", []):
            email = member.get("email")
            sid = member.get("student_id")
            identifier = email or sid or member.get("name")
            if identifier:
                unique_athletes.add(identifier.lower())

            gender = (member.get("gender") or "").lower()
            if "boy" in gender or "male" in gender:
                gender_counts["Boy"] += 1
            elif "girl" in gender or "female" in gender:
                gender_counts["Girl"] += 1
            else:
                gender_counts["Other"] += 1

    return {
        "total_teams": len(teams),
        "total_athletes": len(unique_athletes),
        "dept_counts": dept_counts,
        "sport_counts": sport_counts,
        "category_counts": category_counts,
        "gender_counts": gender_counts,
    }


@router.get("/", response_class=HTMLResponse)
@router.get("/dashboard", response_class=HTMLResponse)
async def dashboard_view(request: Request):
    """Render the comprehensive registered teams viewer dashboard."""
    teams, error = await fetch_all_teams()
    stats = compute_dashboard_stats(teams)

    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "teams": teams,
            "stats": stats,
            "departments": DEPARTMENT_CONFIG,
            "sports": SPORTS_CONFIG,
            "db_error": error,
        },
    )


@router.get("/api/teams", response_class=JSONResponse)
async def api_get_teams(
    department: Optional[str] = Query(None, description="Filter by department code (amsc, ppe, ppt)"),
    sport: Optional[str] = Query(None, description="Filter by sport id"),
    category: Optional[str] = Query(None, description="Filter by category (Boys, Girls, Mixed)"),
    q: Optional[str] = Query(None, description="Search query across team, players, captain, or email"),
):
    """API endpoint to get filtered registered teams in JSON format."""
    teams, error = await fetch_all_teams()
    if error:
        return JSONResponse(status_code=500, content={"error": error, "teams": []})

    filtered = teams
    if department:
        dept_canonical = DEPT_ALIASES.get(department.lower(), department.lower())
        filtered = [t for t in filtered if t.get("department_code") == dept_canonical]

    if sport:
        filtered = [t for t in filtered if t.get("sport_id") == sport]

    if category:
        filtered = [t for t in filtered if t.get("category", "").lower() == category.lower()]

    if q:
        query = q.strip().lower()

        def match_team(t: dict[str, Any]) -> bool:
            if query in t.get("sport_name", "").lower():
                return True
            if query in t.get("department", "").lower():
                return True
            if query in t.get("display_title", "").lower():
                return True
            for m in t.get("all_members", []):
                if query in m.get("name", "").lower():
                    return True
                if query in m.get("email", "").lower():
                    return True
                if query in m.get("student_id", "").lower():
                    return True
            return False

        filtered = [t for t in filtered if match_team(t)]

    return {
        "count": len(filtered),
        "total": len(teams),
        "teams": filtered,
    }


@router.get("/api/teams/{team_id}", response_class=JSONResponse)
async def api_get_team_detail(team_id: str):
    """API endpoint to retrieve single team details by team_id."""
    teams, error = await fetch_all_teams()
    if error:
        raise HTTPException(status_code=500, detail=error)

    for t in teams:
        if t.get("team_id") == team_id:
            return t

    raise HTTPException(status_code=404, detail=f"Team with ID '{team_id}' not found.")


@router.get("/api/stats", response_class=JSONResponse)
async def api_get_stats():
    """API endpoint returning tournament registration metrics."""
    teams, error = await fetch_all_teams()
    stats = compute_dashboard_stats(teams)
    is_connected = await ping_db()

    return {
        "status": "online" if is_connected else "offline",
        "error": error,
        "stats": stats,
    }


def build_sports_meet_excel_workbook(teams: list[dict[str, Any]]) -> io.BytesIO:
    """
    Build a multi-sheet Excel workbook (16 sheets, 1 for each sport).
    Within each sheet, separate registered teams clearly with department-themed banners,
    squad metadata, highlighted captain, zebra-striped player rows, and auto-sized columns.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    if wb.active:
        wb.remove(wb.active)

    # Typography
    font_main_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, bold=True, color="E2E8F0")
    font_summary_sub = Font(name="Segoe UI", size=9, italic=True, color="475569")
    font_th = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_captain = Font(name="Segoe UI", size=10, bold=True, color="92400E")
    font_player = Font(name="Segoe UI", size=10, color="0F172A")
    font_sub = Font(name="Segoe UI", size=10, italic=True, color="78350F")
    font_empty = Font(name="Segoe UI", size=10, italic=True, color="64748B")

    # Fills
    fill_main_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_subtitle = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_th = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_team_subbar = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_captain = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_zebra_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra_2 = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_sub = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
    fill_empty = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    dept_styles = {
        "amsc": {
            "fill": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
            "color": "0369A1",
            "border": Side(style="medium", color="7DD3FC"),
        },
        "ppe": {
            "fill": PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid"),
            "color": "6B21A8",
            "border": Side(style="medium", color="D8B4FE"),
        },
        "ppt": {
            "fill": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
            "color": "065F46",
            "border": Side(style="medium", color="6EE7B7"),
        },
    }

    thin_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    for sport_id, sport_conf in SPORTS_CONFIG.items():
        sheet_title = sport_conf.get("display_name", sport_conf["name"])[:31].replace("/", "-")
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Row 1: Tournament Header Banner
        ws.merge_cells("A1:G1")
        r1 = ws["A1"]
        r1.value = "INTER-DEPARTMENT SPORTS MEET 2026 · IIT ROORKEE"
        r1.font = font_main_title
        r1.fill = fill_main_title
        r1.alignment = align_center
        ws.row_dimensions[1].height = 26

        # Row 2: Sport Category & Squad Rules
        ws.merge_cells("A2:G2")
        r2 = ws["A2"]
        r2.value = f"EVENT: {sport_conf.get('display_name', sport_conf['name'])}   |   CATEGORY: {sport_conf.get('category', 'Open')}   |   SQUAD RULE: {sport_conf.get('squad_notation', '')}"
        r2.font = font_subtitle
        r2.fill = fill_subtitle
        r2.alignment = align_center
        ws.row_dimensions[2].height = 20

        # Row 3: Blank spacing row
        ws.row_dimensions[3].height = 10

        # Teams in this sport
        sport_teams = [t for t in teams if t.get("sport_id") == sport_id]

        if not sport_teams:
            ws.merge_cells("A4:G4")
            c = ws["A4"]
            c.value = "ℹ️ No registered teams yet for this sport event."
            c.font = font_empty
            c.fill = fill_empty
            c.alignment = align_center
            ws.row_dimensions[4].height = 28
            for col in range(1, 8):
                ws.cell(row=4, column=col).border = cell_border
        else:
            curr_row = 4
            for t_idx, team in enumerate(sport_teams, start=1):
                dept_code = (team.get("department_code") or "ppe").lower()
                d_style = dept_styles.get(dept_code, dept_styles["ppe"])

                # 1. Team Header Banner
                ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=7)
                banner_cell = ws.cell(row=curr_row, column=1)
                banner_cell.value = (
                    f"🏆 TEAM #{t_idx}: {team.get('department')} DEPARTMENT  ·  "
                    f"TEAM ID: {team.get('team_id')}  ·  STATUS: {team.get('status', 'Registered').upper()}"
                )
                banner_cell.font = Font(name="Segoe UI", size=11, bold=True, color=d_style["color"])
                banner_cell.fill = d_style["fill"]
                banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                ws.row_dimensions[curr_row].height = 24
                for col in range(1, 8):
                    ws.cell(row=curr_row, column=col).border = Border(top=d_style["border"], bottom=thin_side, left=thin_side, right=thin_side)
                curr_row += 1

                # 2. Team Sub-bar with Squad Metrics & Registration Time
                ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=7)
                sub_cell = ws.cell(row=curr_row, column=1)
                created_by_info = f"Created By: {team.get('created_by_name') or team.get('created_by') or 'N/A'}"
                sub_cell.value = (
                    f"Squad Strength: {team.get('total_strength', 0)} ({team.get('total_playing', 0)} Playing + {team.get('total_substitutes', 0)} Subs)   |   "
                    f"Registered On: {team.get('created_at_formatted', 'N/A')}   |   {created_by_info}"
                )
                sub_cell.font = font_summary_sub
                sub_cell.fill = fill_team_subbar
                sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                ws.row_dimensions[curr_row].height = 18
                for col in range(1, 8):
                    ws.cell(row=curr_row, column=col).border = cell_border
                curr_row += 1

                # 3. Table Column Headers
                headers = [
                    "S.No",
                    "Squad Role",
                    "Athlete Name",
                    "Official IITR Email",
                    "Student ID",
                    "Gender",
                    "Department",
                ]
                for c_idx, h_title in enumerate(headers, start=1):
                    hc = ws.cell(row=curr_row, column=c_idx)
                    hc.value = h_title
                    hc.font = font_th
                    hc.fill = fill_th
                    hc.alignment = align_center if c_idx not in (3, 4) else Alignment(horizontal="left", vertical="center", indent=1)
                    hc.border = cell_border
                ws.row_dimensions[curr_row].height = 22
                curr_row += 1

                # 4. Roster Rows
                p_idx = 1
                captain = team.get("captain") or {}
                row_cells = [
                    (p_idx, align_center),
                    ("👑 Team Captain", align_left),
                    (captain.get("name", "N/A"), align_left),
                    (captain.get("email", "N/A"), align_left),
                    (captain.get("student_id", "N/A"), align_center),
                    (captain.get("gender", "N/A"), align_center),
                    (team.get("department", "N/A"), align_center),
                ]
                for c_idx, (val, al) in enumerate(row_cells, start=1):
                    cell = ws.cell(row=curr_row, column=c_idx)
                    cell.value = val
                    cell.font = font_captain
                    cell.fill = fill_captain
                    cell.alignment = al
                    cell.border = cell_border
                ws.row_dimensions[curr_row].height = 20
                curr_row += 1
                p_idx += 1

                # Main Players
                for mp in team.get("main_players") or []:
                    row_fill = fill_zebra_1 if p_idx % 2 == 0 else fill_zebra_2
                    row_cells = [
                        (p_idx, align_center),
                        (mp.get("role", f"Main Player {p_idx-1}"), align_left),
                        (mp.get("name", "N/A"), align_left),
                        (mp.get("email", "N/A"), align_left),
                        (mp.get("student_id", "N/A"), align_center),
                        (mp.get("gender", "N/A"), align_center),
                        (team.get("department", "N/A"), align_center),
                    ]
                    for c_idx, (val, al) in enumerate(row_cells, start=1):
                        cell = ws.cell(row=curr_row, column=c_idx)
                        cell.value = val
                        cell.font = font_player
                        cell.fill = row_fill
                        cell.alignment = al
                        cell.border = cell_border
                    ws.row_dimensions[curr_row].height = 20
                    curr_row += 1
                    p_idx += 1

                # Substitutes
                for s_i, sub in enumerate(team.get("substitutes") or [], start=1):
                    row_cells = [
                        (p_idx, align_center),
                        (sub.get("role", f"Substitute {s_i}"), align_left),
                        (sub.get("name", "N/A"), align_left),
                        (sub.get("email", "N/A"), align_left),
                        (sub.get("student_id", "N/A"), align_center),
                        (sub.get("gender", "N/A"), align_center),
                        (team.get("department", "N/A"), align_center),
                    ]
                    for c_idx, (val, al) in enumerate(row_cells, start=1):
                        cell = ws.cell(row=curr_row, column=c_idx)
                        cell.value = val
                        cell.font = font_sub
                        cell.fill = fill_sub
                        cell.alignment = al
                        cell.border = cell_border
                    ws.row_dimensions[curr_row].height = 20
                    curr_row += 1
                    p_idx += 1

                # Team Separator Spacing
                ws.row_dimensions[curr_row].height = 12
                curr_row += 1
                ws.row_dimensions[curr_row].height = 8
                curr_row += 1

        # Auto-fit column widths
        min_widths = {1: 8, 2: 20, 3: 28, 4: 32, 5: 16, 6: 12, 7: 15}
        for col_idx in range(1, 8):
            col_letter = get_column_letter(col_idx)
            max_len = min_widths.get(col_idx, 15)
            for row in range(4, ws.max_row + 1):
                cell_val = ws.cell(row=row, column=col_idx).value
                if cell_val and isinstance(cell_val, str) and not ws.cell(row=row, column=col_idx).coordinate in ws.merged_cells:
                    max_len = max(max_len, len(cell_val) + 3)
            ws.column_dimensions[col_letter].width = min(max_len, 42)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@router.get("/api/export/excel")
@router.get("/api/export/workbook")
async def api_export_excel():
    """Export all registered teams into a styled 16-sheet Excel workbook (1 sheet per sport)."""
    teams, error = await fetch_all_teams()
    if error:
        raise HTTPException(status_code=500, detail=error)

    buffer = build_sports_meet_excel_workbook(teams)
    filename = f"inter_dept_sports_meet_rosters_16_sports_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/api/export/csv")
async def api_export_csv(format: Optional[str] = Query(None)):
    """
    Export registered teams.
    Default: Returns the 16-sheet styled workbook (.xlsx) as requested.
    If ?format=raw_csv, returns standard flat CSV file.
    """
    teams, error = await fetch_all_teams()
    if error:
        raise HTTPException(status_code=500, detail=error)

    if format == "raw_csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "Department",
            "Sport",
            "Category",
            "Team ID",
            "Member Role",
            "Member Name",
            "IITR Email",
            "Student ID",
            "Gender",
            "Status",
            "Registration Date",
        ])

        for t in teams:
            dept = t.get("department", "")
            sport = t.get("sport_name", "")
            cat = t.get("category", "")
            tid = t.get("team_id", "")
            status = t.get("status", "Registered")
            reg_date = t.get("created_at_formatted", "")

            for m in t.get("all_members", []):
                writer.writerow([
                    dept,
                    sport,
                    cat,
                    tid,
                    m.get("role", ""),
                    m.get("name", ""),
                    m.get("email", ""),
                    m.get("student_id", ""),
                    m.get("gender", ""),
                    status,
                    reg_date,
                ])

        output.seek(0)
        filename = f"inter_dept_teams_roster_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    # Multi-sheet beautiful workbook
    buffer = build_sports_meet_excel_workbook(teams)
    filename = f"inter_dept_sports_meet_rosters_16_sports_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
